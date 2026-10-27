# -*- coding: utf-8 -*-
"""
تصدير التقارير إلى ملف Excel (.xlsx)
=======================================
دالة عامة واحدة تُستخدم من كل شاشات التقارير (يومي، شهري، أداء كاشيرين،
مخزون، ربح/خسارة، شِفتات) بدل ما نكرر منطق كتابة الإكسل بكل شاشة.

بيستخدم openpyxl - مكتبة خفيفة، بتشتغل عادي على وندوز 7 مع PyInstaller
(pure-python، ما إلها أي اعتمادية بايناري خارجية).
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


def export_rows_to_excel(filepath, sheet_title, headers, rows, summary_lines=None):
    """
    filepath: المسار الكامل اللي بدنا نحفظ فيه الملف (.xlsx)
    sheet_title: اسم الورقة
    headers: قائمة بأسماء الأعمدة
    rows: قائمة من tuples/lists - كل عنصر صف بنفس ترتيب headers
    summary_lines: (اختياري) قائمة من (تسمية, قيمة) تُكتب فوق الجدول (مجاميع، إلخ)
    يرجع (success: bool, message: str)
    """
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title[:31] if sheet_title else "تقرير"
        ws.sheet_view.rightToLeft = True

        start_row = 1
        if summary_lines:
            for label, value in summary_lines:
                ws.cell(row=start_row, column=1, value=label).font = Font(bold=True)
                ws.cell(row=start_row, column=2, value=value)
                start_row += 1
            start_row += 1  # سطر فاضي قبل الجدول

        header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_offset, row_data in enumerate(rows, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row=start_row + row_offset, column=col_idx, value=value)

        # عرض أعمدة تلقائي تقريبي حسب أطول محتوى
        for col_idx, header in enumerate(headers, start=1):
            max_len = len(str(header))
            for row_data in rows:
                if col_idx - 1 < len(row_data):
                    max_len = max(max_len, len(str(row_data[col_idx - 1])))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 4, 45)

        wb.save(filepath)
        return True, f"تم تصدير التقرير بنجاح إلى:\n{filepath}"
    except Exception as e:
        return False, f"تعذّر تصدير التقرير: {e}"
